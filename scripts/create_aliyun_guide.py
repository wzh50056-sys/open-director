from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/阿里云上线资源采购与部署指南.docx")
BLUE = "1769AA"
DARK = "17324D"
LIGHT = "EAF3F9"
PALE = "F5F8FA"
GRAY = "667085"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=11, bold=False, color=None, name="PingFang SC"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "PingFang SC")
    fonts.set(qn("w:ascii"), "PingFang SC")
    fonts.set(qn("w:hAnsi"), "PingFang SC")
    r_pr.append(fonts)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, size=9, color=GRAY)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_step(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    p = cell.paragraphs[0]
    r = p.add_run(title + "　")
    set_font(r, bold=True, color=DARK)
    r = p.add_run(text)
    set_font(r, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_link_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["资源", "用途", "阿里云官方入口"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, bold=True, color=WHITE)
    for resource, purpose, label, url in rows:
        cells = table.add_row().cells
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cells[0].paragraphs[0].add_run(resource)
        set_font(r, bold=True, color=DARK)
        r = cells[1].paragraphs[0].add_run(purpose)
        set_font(r)
        add_hyperlink(cells[2].paragraphs[0], label, url)
    set_table_geometry(table, [1850, 3200, 4310])
    return table


def add_config_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, text in enumerate(("配置项", "推荐选择")):
        set_cell_shading(table.rows[0].cells[i], BLUE)
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, bold=True, color=WHITE)
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE)
        r = cells[0].paragraphs[0].add_run(label)
        set_font(r, bold=True, color=DARK)
        r = cells[1].paragraphs[0].add_run(value)
        set_font(r)
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_geometry(table, [2700, 6660])
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.4)
section.footer_distance = Inches(0.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "PingFang SC"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.22

for level, size, before, after, color in ((1, 16, 18, 9, BLUE), (2, 13, 13, 6, BLUE), (3, 11.5, 9, 4, DARK)):
    style = styles[f"Heading {level}"]
    style.font.name = "PingFang SC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Number"):
    style = styles[style_name]
    style.font.name = "PingFang SC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.2

header_p = section.header.paragraphs[0]
header_p.text = "OPEN DIRECTOR　｜　阿里云生产环境部署指南"
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in header_p.runs:
    set_font(run, size=8.5, bold=True, color=GRAY)
add_page_number(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(54)
p.paragraph_format.space_after = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("阿里云上线资源采购与部署指南")
set_font(r, size=25, bold=True, color=DARK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
r = p.add_run("适用于 OpenDirector：Web、Canvas 与 Worker 同机起步方案")
set_font(r, size=13, color=BLUE)

add_note(doc, "推荐起步目标", "日活约 1 万人、峰值在线约 300～800 人；Web 与 Worker 先部署在一台 ECS，RDS、Tair 和 OSS 使用阿里云托管服务。")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("推荐架构")
set_font(r, size=12, bold=True, color=DARK)

architecture = (
    "用户\n"
    "  ↓\n"
    "CDN + HTTPS\n"
    "  ↓\n"
    "ECS 8核16GB\n"
    "├── Nginx\n"
    "├── Web 容器\n"
    "├── Canvas 画布服务\n"
    "└── Worker 容器\n"
    "       ↓\n"
    "  RDS MySQL  ·  Tair Redis  ·  OSS 图片/音视频"
)
table = doc.add_table(rows=1, cols=1)
set_table_geometry(table, [CONTENT_DXA])
cell = table.cell(0, 0)
set_cell_shading(cell, "F0F5F9")
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(architecture)
set_font(r, size=10.5, bold=True, color=DARK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(16)
r = p.add_run("版本：2026年8月　｜　采购前请以阿里云控制台实时库存与价格为准")
set_font(r, size=9, color=GRAY)

doc.add_page_break()

add_heading(doc, "1. 官方购买与开通入口", 1)
p = doc.add_paragraph("下表均为阿里云官方入口。ECS、RDS 和 Tair需要创建实例；OSS、CDN和证书服务通常先开通，再按实际使用或套餐计费。")

links = [
    ("阿里云账号", "注册、登录及实名认证", "登录/注册", "https://account.aliyun.com/"),
    ("域名", "购买网站域名", "万网域名", "https://wanwang.aliyun.com/domain/"),
    ("ECS", "运行Nginx、Web、Canvas和Worker", "云服务器ECS", "https://cn.aliyun.com/product/ecs"),
    ("RDS MySQL", "保存用户、项目、任务、订单和权限数据", "RDS MySQL", "https://cn.aliyun.com/product/rds/mysql"),
    ("Tair/Redis", "缓存、Session、限流和BullMQ任务队列", "Tair控制台", "https://kvstore.console.aliyun.com/"),
    ("OSS", "保存图片、音频、视频和用户上传文件", "对象存储OSS", "https://cn.aliyun.com/product/oss"),
    ("CDN", "加速OSS中的图片和视频", "内容分发网络CDN", "https://cn.aliyun.com/product/cdn"),
    ("HTTPS证书", "为网站、API和CDN域名启用HTTPS", "数字证书管理", "https://cn.aliyun.com/product/cas"),
    ("DNS", "把域名解析到ECS或CDN", "云解析DNS", "https://dns.console.aliyun.com/"),
    ("ICP备案", "中国大陆服务器正式开站备案", "阿里云备案", "https://beian.aliyun.com/"),
    ("VPC", "让ECS、RDS和Tair通过私网通信", "VPC控制台", "https://vpc.console.aliyun.com/"),
    ("费用中心", "预算、账单和费用告警", "费用与成本", "https://billing-cost.console.aliyun.com/"),
]
add_link_table(doc, links)

add_note(doc, "关键原则", "ECS、RDS、Tair和OSS应选择同一地域；ECS、RDS和Tair应加入同一VPC，应用使用内网地址连接数据库和Redis。")

add_heading(doc, "2. ECS 服务器配置", 1)
add_config_table(doc, [
    ("用途", "同机运行 Nginx、Web、Canvas 和 Worker"),
    ("实例", "通用型 g 系列或计算型 c 系列，x86 架构"),
    ("起步规格", "8核16GB；视频渲染较重时优先16核32GB"),
    ("操作系统", "Ubuntu 22.04 LTS 64位"),
    ("系统盘", "100GB ESSD"),
    ("数据盘", "建议200GB ESSD，用于渲染临时文件"),
    ("公网带宽", "5～10Mbps起步；图片和视频下载交给OSS + CDN"),
    ("安全组", "开放80/443；22端口限制为管理员IP；不要开放3306/6379"),
])

add_heading(doc, "容器资源建议", 2)
add_config_table(doc, [
    ("Web容器", "约3核、5GB内存，负责页面和API"),
    ("Canvas服务", "约1核、2GB内存，负责画布后端"),
    ("Worker容器", "约3核、7GB内存，WORKER_CONCURRENCY=1"),
    ("系统保留", "至少保留1核和2GB给Ubuntu、Nginx与Docker"),
])

add_heading(doc, "3. RDS MySQL配置", 1)
add_config_table(doc, [
    ("版本", "MySQL 8.0"),
    ("系列", "高可用系列，一主一备"),
    ("起步规格", "2核4GB；业务增长后升级4核8GB"),
    ("存储", "100GB ESSD PL1，开启存储自动扩容"),
    ("网络", "VPC内网，不申请公网连接地址"),
    ("备份", "每日自动备份，保留7～14天"),
    ("安全", "白名单只允许ECS所在交换机或安全组访问"),
])
add_note(doc, "不要选择", "正式生产环境不要选择基础单节点系列；不要把RDS的3306端口开放到公网。")

add_heading(doc, "4. Tair / Redis配置", 1)
add_config_table(doc, [
    ("产品", "Tair（兼容Redis）"),
    ("架构", "标准主备双副本，初期不使用单节点"),
    ("容量", "2GB起步，推荐4GB"),
    ("用途", "BullMQ任务、Session、缓存、验证码和限流"),
    ("网络", "与ECS相同VPC，只使用内网连接"),
    ("备份", "开启持久化、自动备份和内存告警"),
])
add_bullet(doc, "已完成和失败的BullMQ任务必须设置保留时长与数量上限。")
add_bullet(doc, "Redis不保存唯一一份订单、用户或项目数据；永久数据写入RDS。")
add_bullet(doc, "Redis内存持续超过60%～70%时，升级到4GB或8GB。")

add_heading(doc, "5. OSS文件存储配置", 1)
add_config_table(doc, [
    ("Bucket权限", "私有；开启阻止公共访问"),
    ("存储类型", "标准存储"),
    ("冗余", "预算允许选择同城冗余ZRS"),
    ("地域", "与ECS相同"),
    ("安全", "服务端加密；使用RAM子账号，不使用主账号AccessKey"),
    ("访问", "后端生成签名URL，或由CDN提供鉴权URL"),
])

add_heading(doc, "建议目录", 2)
for text in ("users/：用户相关文件", "uploads/：用户上传", "images/：生成图片", "audio/：语音与音乐", "videos/：视频素材", "renders/：最终成片", "temporary/：临时文件"):
    add_bullet(doc, text)

add_heading(doc, "生命周期规则", 2)
add_config_table(doc, [
    ("临时上传/中间文件", "1～7天自动删除"),
    ("失败任务文件", "7～30天自动删除"),
    ("原始素材", "30～90天后按访问情况转低频存储"),
    ("长期未访问成片", "90天后转低频或归档；注意最低保存周期和取回费用"),
])

add_heading(doc, "6. CDN、域名与HTTPS", 1)
add_config_table(doc, [
    ("网站域名", "www.example.com → ECS公网IP"),
    ("素材域名", "cdn.example.com → CDN提供的CNAME"),
    ("CDN源站", "OSS Bucket"),
    ("业务类型", "图片小文件或视音频点播"),
    ("HTTPS", "网站与CDN均启用；证书可使用阿里云免费DV证书"),
    ("安全", "开启URL鉴权、Referer防盗链、带宽封顶和费用告警"),
    ("缓存", "图片和最终视频通常缓存7～30天；API不做长缓存"),
])

add_note(doc, "备案提示", "ECS或CDN使用中国大陆地域时，域名通常需要完成ICP备案；香港地域一般不要求ICP备案，但中国大陆访问质量可能有所差异。")

add_heading(doc, "7. 推荐购买与部署顺序", 1)
steps = [
    "注册阿里云账号并完成个人或企业实名认证。",
    "购买并实名认证域名。",
    "购买ECS，创建VPC、交换机和安全组。",
    "若使用中国大陆ECS，立即提交ICP备案。",
    "创建RDS MySQL高可用实例，确保地域和VPC与ECS一致。",
    "创建Tair标准主备实例，确保地域和VPC与ECS一致。",
    "开通OSS并创建私有Bucket，配置RAM权限和生命周期。",
    "申请HTTPS证书，配置Nginx和正式域名。",
    "部署Web、Canvas和Worker容器并完成数据库迁移。",
    "确认OSS上传下载正常后开通CDN，并配置CNAME、鉴权和缓存。",
    "设置CPU、内存、磁盘、数据库、Redis和费用告警。",
    "进行压力测试和视频任务测试，通过后正式开放流量。",
]
for step in steps:
    add_step(doc, step)

add_heading(doc, "8. 上线前检查清单", 1)
checks = [
    "ECS、RDS、Tair和OSS地域一致；ECS、RDS和Tair位于同一VPC。",
    "MySQL和Redis均未开放公网访问。",
    "OSS Bucket为私有，并开启阻止公共访问。",
    "Web、Canvas和Worker容器配置自动重启与日志轮转。",
    "Worker并发初始设置为1，且不会抢占全部CPU和内存。",
    "渲染成功、失败和超时后都会清理临时文件。",
    "RDS每日备份，并进行过至少一次恢复演练。",
    "CDN开启HTTPS、URL鉴权、防盗链和带宽封顶。",
    "配置API限流、单用户任务上限和AI费用预算。",
    "完成注册、登录、上传、画布、生成、渲染和播放全链路测试。",
]
for check in checks:
    add_bullet(doc, "□ " + check)

add_heading(doc, "9. 后期扩容路线", 1)
add_config_table(doc, [
    ("阶段A", "ECS 8核16GB同机运行Web + Canvas + Worker；Worker并发1"),
    ("阶段B", "当渲染影响网站或任务等待过长时，把Worker拆到独立8核16GB ECS"),
    ("阶段C", "网站需要高可用时，增加第二台Web ECS并接入ALB"),
    ("阶段D", "根据Web RPS和队列长度分别自动扩容Web与Worker"),
    ("阶段E", "日活达到数万后，再评估Tair集群、RDS只读实例、ACK或更细服务化"),
])

add_note(doc, "容量说明", "8核16GB同机方案可以作为日活约5,000～10,000的起步目标，但实际容量取决于接口、缓存、数据库查询和视频任务耗时。正式承诺容量前必须进行压测。")

add_heading(doc, "10. 关键官方资料", 1)
sources = [
    ("ECS入门指引", "https://help.aliyun.com/zh/ecs/quick-start"),
    ("RDS MySQL产品系列", "https://help.aliyun.com/zh/rds/apsaradb-rds-for-mysql/product-editions/"),
    ("Tair产品架构", "https://help.aliyun.com/zh/redis/product-overview/product-architecture"),
    ("OSS产品与安全能力", "https://cn.aliyun.com/product/oss"),
    ("CDN产品说明", "https://cn.aliyun.com/product/cdn"),
    ("阿里云性能测试PTS", "https://help.aliyun.com/zh/pts/"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    add_hyperlink(p, label, url)

# Avoid isolated headings and repeat table headers.
for table in doc.tables:
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT.resolve())
